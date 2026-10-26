from pptx import Presentation

prs = Presentation('output/consolidated_mbr.pptx')

print(f"Total slides: {len(prs.slides)}\n")

for i, slide in enumerate(prs.slides, 1):
    title = slide.shapes.title.text if slide.shapes.title else "No title"
    shapes_count = len(slide.shapes)
    text_shapes = sum(1 for shape in slide.shapes if shape.has_text_frame)
    print(f"Slide {i}: {title}")
    print(f"   Shapes: {shapes_count} total, {text_shapes} with text")
