"""
Create March MBR by combining template with team presentation slides.
Uses proper slide ordering.
"""
from pptx import Presentation
from pathlib import Path
from copy import deepcopy
import sys

def copy_slide(source_prs, slide_idx, dest_prs):
    """Copy slide from source to destination."""
    try:
        source_slide = source_prs.slides[slide_idx - 1]
        blank_layout = dest_prs.slide_layouts[6]
        dest_slide = dest_prs.slides.add_slide(blank_layout)
        
        for shape in source_slide.shapes:
            el = shape.element
            newel = deepcopy(el)
            dest_slide.shapes._spTree.insert_element_before(newel, 'p:extLst')
        
        return True
    except Exception as e:
        print(f"      ⚠️ Error: {e}")
        return False

def find_slide_idx(prs, keyword):
    """Find slide by keyword in title."""
    for i, slide in enumerate(prs.slides, 1):
        if slide.shapes.title and keyword.lower() in slide.shapes.title.text.lower():
            return i
    return None

def main():
    print("\n" + "="*80)
    print("🎯 CREATING MARCH 2026 MBR REPORT")
    print("="*80 + "\n")
    
    template_path = "sample_data/final-data/06-01-2026.pptx"
    ppt_folder = "sample_data/ppt-data"
    output_path = "output/consolidated_mbr.pptx"
    
    # Load template
    print("📋 Step 1: Loading template...")
    template = Presentation(template_path)
    print(f"   Template: {len(template.slides)} slides\n")
    
    # Create new presentation with template slides 1-5
    print("📋 Step 2: Creating new presentation...")
    new_prs = Presentation()
    
    # Remove default slide
    if len(new_prs.slides) > 0:
        rId = new_prs.slides._sldIdLst[0].rId
        new_prs.part.drop_rel(rId)
        del new_prs.slides._sldIdLst[0]
    
    # Copy template slides 1-5 (overview)
    print("   Copying template overview slides (1-5)...")
    for i in range(1, 6):
        copy_slide(template, i, new_prs)
        print(f"      ✓ Slide {i}: {template.slides[i-1].shapes.title.text if template.slides[i-1].shapes.title else 'No title'}")
    
    # Find and add March team slides
    print("\n📋 Step 3: Adding March team slides...")
    
    ppt_path = Path(ppt_folder)
    team_files = sorted(list(ppt_path.glob("*March*.pptx")))
    team_files = [f for f in team_files if not f.name.startswith('~$')]
    
    for ppt_file in team_files:
        team = ppt_file.stem.replace("V2Solutions_", "").replace("LT_", "").replace("_March_2026_Report", "").replace("_March2026_Report", "").replace("_March_Report", "").replace("_Mar_2026_Report 1", "")
        
        print(f"\n   📄 {team}")
        
        try:
            source = Presentation(str(ppt_file))
            
            # Goals Achieved
            goals_idx = find_slide_idx(source, "Goals Achieved")
            if goals_idx:
                if copy_slide(source, goals_idx, new_prs):
                    print(f"      ✓ Goals Achieved (Slide {goals_idx})")
            
            # In Progress
            progress_idx = find_slide_idx(source, "In Progress")
            if progress_idx:
                if copy_slide(source, progress_idx, new_prs):
                    print(f"      ✓ In Progress (Slide {progress_idx})")
                    
        except Exception as e:
            print(f"      ⚠️ Error: {e}")
    
    # Copy template slides 6-14 (Project Details, Metrics, Thank You)
    print("\n📋 Step 4: Adding template ending slides (6-14)...")
    for i in range(6, 15):
        if i <= len(template.slides):
            copy_slide(template, i, new_prs)
            title = template.slides[i-1].shapes.title.text if template.slides[i-1].shapes.title else "No title"
            print(f"      ✓ Slide {i}: {title}")
    
    # Save
    print(f"\n💾 Step 5: Saving...")
    new_prs.save(output_path)
    
    print("\n" + "="*80)
    print("✅ SUCCESS!")
    print("="*80)
    print(f"\n📁 Saved: {output_path}")
    print(f"📊 Total slides: {len(new_prs.slides)}")
    print(f"\nStructure:")
    print(f"  Slides 1-5: Template Overview")
    print(f"  Slides 6-X: March Goals & In-Progress (all teams)")
    print(f"  Slides X+: Project Metrics & Thank You")
    print()

if __name__ == "__main__":
    main()
