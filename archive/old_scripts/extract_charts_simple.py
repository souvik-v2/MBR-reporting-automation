"""
Simple approach: Copy charts manually OR use screenshot approach.

Since COM automation is complex, let's document the manual process
and create placeholder slides that can be easily updated.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pathlib import Path


def create_project_metrics_slide(prs, project_name, achievements, layout_idx=6):
    """
    Create a Project Metrics slide with placeholders for charts.
    
    Args:
        prs: Presentation object
        project_name: Name of the project
        achievements: List of achievement strings
        layout_idx: Layout index (6 = blank)
    
    Returns:
        slide object
    """
    # Use blank layout
    blank_layout = prs.slide_layouts[layout_idx]
    slide = prs.slides.add_slide(blank_layout)
    
    # Add title
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.3),
        Inches(9), Inches(0.6)
    )
    title_frame = title_box.text_frame
    title_frame.text = f"Project Metrics - {project_name}"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    
    # Add placeholder for Rework Trend chart (left)
    chart_placeholder1 = slide.shapes.add_textbox(
        Inches(0.5), Inches(1.2),
        Inches(4.5), Inches(3)
    )
    chart_frame1 = chart_placeholder1.text_frame
    chart_frame1.text = "[Rework Trend Chart]\n\nPaste chart from Excel here:\n• Open LT-SimplifiedMetrics-{project}.xlsx\n• Copy 'Rework Trend' chart\n• Paste here"
    chart_frame1.paragraphs[0].font.size = Pt(14)
    chart_frame1.paragraphs[0].font.italic = True
    
    # Add placeholder for Throughput Distribution chart (right)
    chart_placeholder2 = slide.shapes.add_textbox(
        Inches(5.2), Inches(1.2),
        Inches(4.5), Inches(3)
    )
    chart_frame2 = chart_placeholder2.text_frame
    chart_frame2.text = "[Throughput Distribution Chart]\n\nPaste chart from Excel here:\n• Copy 'Throughput Distribution' chart\n• Paste here"
    chart_frame2.paragraphs[0].font.size = Pt(14)
    chart_frame2.paragraphs[0].font.italic = True
    
    # Add Key Achievements section
    achievements_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(4.5),
        Inches(9), Inches(2.5)
    )
    achievements_frame = achievements_box.text_frame
    achievements_frame.word_wrap = True
    
    # Title
    p = achievements_frame.paragraphs[0]
    p.text = "Key Achievements:"
    p.font.size = Pt(18)
    p.font.bold = True
    
    # Achievements list
    for achievement in achievements[:6]:  # Max 6 achievements
        p = achievements_frame.add_paragraph()
        p.text = f"• {achievement}"
        p.font.size = Pt(14)
        p.level = 0
    
    return slide


def test_create_slide():
    """Test creating a metrics slide."""
    print("\n" + "="*80)
    print("🧪 TESTING PROJECT METRICS SLIDE CREATION")
    print("="*80 + "\n")
    
    # Load template
    template = Presentation("sample_data/final-data/06-01-2026.pptx")
    
    # Sample achievements
    achievements = [
        "Typography & UI Enhancements",
        "Mobile Rendering & Responsiveness Fixes",
        "Shortcode Development & Fixes",
        "Accessibility (A11y) Improvements",
        "Table UX & Interaction Improvements",
        "SEO & Indexing Improvements"
    ]
    
    # Create slide
    print("Creating Project Metrics slide for SEOD...")
    slide = create_project_metrics_slide(template, "SEOD", achievements)
    print("   ✓ Slide created\n")
    
    # Save
    output_path = "output/test_metrics_slide.pptx"
    template.save(output_path)
    
    print("="*80)
    print("✅ TEST COMPLETE")
    print("="*80)
    print(f"\n📁 Saved: {output_path}")
    print(f"📊 Total slides: {len(template.slides)}")
    print(f"\nLast slide is 'Project Metrics - SEOD' with:")
    print("  • Title")
    print("  • Placeholder for Rework Trend chart")
    print("  • Placeholder for Throughput Distribution chart")
    print("  • Key Achievements list\n")


if __name__ == "__main__":
    test_create_slide()
