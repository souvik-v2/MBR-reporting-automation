"""Check what's in slide 3"""
from pptx import Presentation

prs = Presentation("output/consolidated_mbr.pptx")
slide3 = prs.slides[2]

print("Slide 3 content:\n")

for shape in slide3.shapes:
    if shape.has_table:
        table = shape.table
        print(f"Table: {len(table.rows)} rows x {len(table.columns)} columns\n")
        
        for row_idx in range(min(7, len(table.rows))):
            row_text = " | ".join([table.cell(row_idx, col_idx).text[:30] for col_idx in range(len(table.columns))])
            print(f"Row {row_idx}: {row_text}")
