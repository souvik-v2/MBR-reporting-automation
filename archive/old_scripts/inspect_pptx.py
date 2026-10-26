"""
Inspect PowerPoint files to see their structure
"""
from pptx import Presentation
from pathlib import Path

ppt_folder = Path("sample_data/ppt-data")
ppt_files = list(ppt_folder.glob("*.pptx"))
ppt_files = [p for p in ppt_files if not p.name.startswith('~$')]

print("\n" + "="*80)
print("INSPECTING POWERPOINT FILES")
print("="*80 + "\n")

for ppt_file in sorted(ppt_files)[:2]:  # Check first 2 files
    print(f"📄 {ppt_file.name}")
    print("-" * 80)
    
    try:
        prs = Presentation(str(ppt_file))
        print(f"Total slides: {len(prs.slides)}\n")
        
        for i, slide in enumerate(prs.slides, 1):
            # Get title
            title = "No title"
            if slide.shapes.title:
                title = slide.shapes.title.text
            
            print(f"Slide {i}: {title}")
            
            # Count shapes
            text_shapes = 0
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text_shapes += 1
            
            print(f"  - Shapes: {len(slide.shapes)} total, {text_shapes} with text")
            
            # Show some content
            if slide.shapes.title:
                print(f"  - Title: {slide.shapes.title.text[:50]}...")
            
            print()
        
    except Exception as e:
        print(f"ERROR: {e}\n")
    
    print("="*80 + "\n")
