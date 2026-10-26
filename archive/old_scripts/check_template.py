from pptx import Presentation

prs = Presentation('sample_data/final-data/06-01-2026.pptx')

print(f"Template total slides: {len(prs.slides)}\n")

for i, slide in enumerate(prs.slides, 1):
    title = slide.shapes.title.text if slide.shapes.title else "No title"
    print(f"Slide {i}: {title}")
