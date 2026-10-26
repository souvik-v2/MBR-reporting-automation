"""
Find each team's real "Project Roadmap" slide and extract its bullets.

Reuses key_highlights_extractor.py's generic slide-bullet extraction
(noise filtering for bare slide numbers, etc.) rather than duplicating
that logic -- this module only adds the title-matching pattern for
"Project Roadmap" / "Project RoadMap" (both capitalizations appear in
the real source decks) and a higher bullet cap, since roadmap lists are
often longer than the 3-bullet cap used for the Key Highlights table.
"""
from __future__ import annotations

import re
from dataclasses import dataclass

from .key_highlights_extractor import extract_summary_from_slide, ExtractiveSummary

ROADMAP_PATTERN = re.compile(r"project\s*road\s*map", re.IGNORECASE)
MAX_ROADMAP_BULLETS = 15


@dataclass
class RoadmapInfo:
    slide_index: int
    title: str
    bullets: list[str]


def find_roadmap_summary(src_prs, max_bullets: int = MAX_ROADMAP_BULLETS) -> "RoadmapInfo | None":
    for i, slide in enumerate(src_prs.slides):
        title = (slide.shapes.title.text if slide.shapes.title else "").strip()
        if ROADMAP_PATTERN.search(title):
            summary = extract_summary_from_slide(slide, max_bullets=max_bullets)
            if summary is not None:
                return RoadmapInfo(slide_index=i, title=title, bullets=summary.bullets)
    return None


if __name__ == "__main__":
    import sys
    from pptx import Presentation

    deck_path = sys.argv[1] if len(sys.argv) > 1 else "sample_data/ppt-data/V2Solutions_SEOD_March_2026_Report.pptx"
    prs = Presentation(deck_path)
    info = find_roadmap_summary(prs)
    if info is None:
        print("No Project Roadmap slide found")
    else:
        print(f"Slide {info.slide_index + 1}: {info.title!r}")
        for b in info.bullets:
            print(" -", b)
