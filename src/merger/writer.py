"""
Writes the final master .xlsx workbook: Master, All_Records, PPTX_Text,
PPTX_Charts, and Run_Log sheets — formatted with a consistent professional
style.
"""
from __future__ import annotations

from typing import List

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1E2761", end_color="1E2761")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT_NAME)


def _sanitize_cell_value(val):
    """Strip control characters openpyxl refuses to write (e.g. the
    vertical-tab \\x0b PowerPoint uses internally as a soft line-break
    inside text boxes -- common in PPTX-sourced text/table rows).
    Replaces them with a space rather than dropping them outright, so
    e.g. "Dayanand\\x0bHarpreet" reads as "Dayanand Harpreet" instead of
    "DayanandHarpreet"."""
    if isinstance(val, str):
        return ILLEGAL_CHARACTERS_RE.sub(" ", val)
    return val


def order_columns(records: List[dict], priority: List[str]) -> List[str]:
    all_cols = []
    seen = set()
    for col in priority:
        if any(col in r for r in records) and col not in seen and not col.startswith("__"):
            all_cols.append(col)
            seen.add(col)
    for r in records:
        for col in r.keys():
            if col not in seen and not col.startswith("__"):
                all_cols.append(col)
                seen.add(col)
    trailing = [c for c in ("__source_file", "__source_sheet") if any(c in r for r in records)]
    return all_cols + trailing


def _write_sheet(ws: Worksheet, records: List[dict], priority_cols: List[str]):
    if not records:
        ws.append(["No data found"])
        ws["A1"].font = BODY_FONT
        return

    columns = order_columns(records, priority_cols)
    display_headers = [c.replace("__", "").replace("_", " ").strip().title() if c.startswith("__") else c for c in columns]

    ws.append([_sanitize_cell_value(h) for h in display_headers])
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

    for row in records:
        out_row = []
        for col in columns:
            val = row.get(col)
            if val is None:
                val = ""
            out_row.append(_sanitize_cell_value(val))
        ws.append(out_row)

    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.font = BODY_FONT

    for idx, col in enumerate(columns, start=1):
        max_len = max(
            [len(str(col))] + [len(str(r.get(col, ""))) for r in records[:500]]
        )
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 10), 50)

    ws.auto_filter.ref = ws.dimensions


def write_master_workbook(
    output_path: str,
    master_records: List[dict],
    all_records: List[dict],
    pptx_text_rows: List[dict],
    pptx_chart_rows: List[dict],
    log: List[str],
    id_col: str,
    date_col: str,
):
    wb = Workbook()

    ws_master = wb.active
    ws_master.title = "Master"
    _write_sheet(ws_master, master_records, [id_col, date_col])

    ws_all = wb.create_sheet("All_Records")
    _write_sheet(ws_all, all_records, [id_col, date_col])

    ws_text = wb.create_sheet("PPTX_Text")
    _write_sheet(ws_text, pptx_text_rows, ["Slide", "Placeholder", "Text"])

    ws_charts = wb.create_sheet("PPTX_Charts")
    _write_sheet(ws_charts, pptx_chart_rows, ["Slide", "Chart Title", "Series", "Category", "Value"])

    ws_log = wb.create_sheet("Run_Log")
    ws_log.append(["Log message"])
    ws_log["A1"].font = HEADER_FONT
    ws_log["A1"].fill = HEADER_FILL
    for msg in log:
        ws_log.append([_sanitize_cell_value(msg)])
    for row_cells in ws_log.iter_rows(min_row=2):
        for cell in row_cells:
            cell.font = BODY_FONT
    ws_log.column_dimensions["A"].width = 110

    wb.save(output_path)
