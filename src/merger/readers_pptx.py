"""
Reader for .pptx files: extracts tables (as row records), free text
(titles/body placeholders), and chart data (categories + series values).
"""
from __future__ import annotations

from pathlib import Path
from typing import List, Tuple

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .column_utils import canonicalize_columns

from . import config


def _table_to_records(table, path_name: str, slide_idx: int, shape_idx: int) -> List[dict]:
    rows = list(table.rows)
    if len(rows) < 2:
        return []
    header_cells = [cell.text.strip() for cell in rows[0].cells]
    headers = canonicalize_columns(header_cells)
    records = []
    for r in rows[1:]:
        values = [cell.text.strip() for cell in r.cells]
        if not any(values):
            continue
        row = dict(zip(headers, values))
        row["__source_file"] = path_name
        row["__source_sheet"] = f"slide{slide_idx}_table{shape_idx}"
        records.append(row)
    return records


def _extract_text_from_shape(shape, path_name: str, slide_idx: int, text_rows: list):
    if not shape.has_text_frame:
        return
    text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
    if not text.strip():
        return
    placeholder_type = "title" if (shape.is_placeholder and shape.placeholder_format.idx == 0) else "body"
    try:
        if shape.is_placeholder and shape.placeholder_format.type is not None:
            placeholder_type = str(shape.placeholder_format.type).split(".")[-1].split(" ")[0].lower()
    except Exception:
        pass
    text_rows.append({
        "__source_file": path_name,
        "Slide": slide_idx,
        "Placeholder": placeholder_type,
        "Text": text.strip(),
    })


def _extract_chart_data(shape, path_name: str, slide_idx: int, shape_idx: int, chart_rows: list):
    try:
        chart = shape.chart
    except Exception:
        return
    try:
        title = chart.chart_title.text_frame.text if chart.has_title else f"Chart {shape_idx}"
    except Exception:
        title = f"Chart {shape_idx}"

    try:
        categories = list(chart.plots[0].categories) if chart.plots else []
    except Exception:
        categories = []

    try:
        for plot in chart.plots:
            for series in plot.series:
                values = list(series.values)
                for i, val in enumerate(values):
                    cat = categories[i] if i < len(categories) else f"item {i+1}"
                    chart_rows.append({
                        "__source_file": path_name,
                        "Slide": slide_idx,
                        "Chart Title": str(title).strip() or f"Chart {shape_idx}",
                        "Series": series.name,
                        "Category": str(cat),
                        "Value": val,
                    })
    except Exception:
        pass


def read_pptx_file(path: Path, log: list) -> Tuple[List[dict], List[dict], List[dict]]:
    """Returns (table_records, text_rows, chart_rows)."""
    table_records, text_rows, chart_rows = [], [], []
    try:
        prs = Presentation(str(path))
    except Exception as exc:
        log.append(f"[ERROR] Could not read PPTX file '{path.name}': {exc}")
        return [], [], []

    for slide_idx, slide in enumerate(prs.slides, start=1):
        for shape_idx, shape in enumerate(slide.shapes, start=1):
            try:
                if config.EXTRACT_PPTX_TABLES and shape.has_table:
                    table_records.extend(
                        _table_to_records(shape.table, path.name, slide_idx, shape_idx)
                    )
                elif config.EXTRACT_PPTX_CHARTS and shape.shape_type == MSO_SHAPE_TYPE.CHART:
                    _extract_chart_data(shape, path.name, slide_idx, shape_idx, chart_rows)
                elif config.EXTRACT_PPTX_TEXT:
                    _extract_text_from_shape(shape, path.name, slide_idx, text_rows)
            except Exception as exc:
                log.append(
                    f"[WARN] Skipped a shape on slide {slide_idx} of '{path.name}': {exc}"
                )

    if not table_records and not text_rows and not chart_rows:
        log.append(f"[INFO] No usable content extracted from '{path.name}'.")

    return table_records, text_rows, chart_rows
