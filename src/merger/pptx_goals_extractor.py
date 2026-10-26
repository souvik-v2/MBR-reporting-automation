"""
Extract 'Goals Achieved' slides content and create a summary report.
Since python-pptx has limitations with slide copying, this creates:
1. A detailed Excel report with all goals content
2. Instructions for manual slide consolidation
"""
from __future__ import annotations
from typing import List, Dict, Any
from pptx import Presentation
from pathlib import Path
import pandas as pd


def extract_slide_content(slide) -> Dict[str, Any]:
    """Extract all text content from a slide."""
    content = {
        'title': '',
        'text_content': [],
        'table_data': [],
        'has_images': False,
        'has_charts': False,
        'shape_count': len(slide.shapes)
    }
    
    for shape in slide.shapes:
        # Get title
        if shape ==slide.shapes.title and shape.has_text_frame:
            content['title'] = shape.text
        
        # Get text content
        elif shape.has_text_frame:
            text = shape.text.strip()
            if text:
                content['text_content'].append(text)
        
        # Get table data
        elif shape.has_table:
            table = shape.table
            table_rows = []
            for row in table.rows:
                row_data = [cell.text for cell in row.cells]
                table_rows.append(row_data)
            content['table_data'].append(table_rows)
        
        # Check for images/charts
        elif shape.shape_type == 13:  # Picture
            content['has_images'] = True
        elif shape.shape_type == 3:  # Chart
            content['has_charts'] = True
    
    return content


def find_and_extract_goals(ppt_path: str, month: str = "May") -> List[Dict[str, Any]]:
    """Find and extract 'Goals Achieved' slides."""
    prs = Presentation(ppt_path)
    goals_data = []
    
    patterns = [
        f"Goals Achieved in {month}",
        f"Goals Achieved",
        f"{month} Goals",
        f"{month} Achievements",
    ]
    
    for i, slide in enumerate(prs.slides):
        if slide.shapes.title:
            title = slide.shapes.title.text
            for pattern in patterns:
                if pattern.lower() in title.lower():
                    content = extract_slide_content(slide)
                    content['source_file'] = Path(ppt_path).name
                    content['slide_number'] = i + 1
                    content['month'] = month
                    goals_data.append(content)
                    break
    
    return goals_data


def create_goals_report(
    ppt_data_folder: str,
    output_excel: str,
    output_instructions: str,
    month: str = "May"
) -> Dict[str, Any]:
    """
    Create comprehensive report of goals slides.
    
    Outputs:
    - Excel file with all goals content
    - Text file with manual consolidation instructions
    """
    ppt_folder = Path(ppt_data_folder)
    team_presentations = list(ppt_folder.glob("*.pptx"))
    team_presentations = [p for p in team_presentations if not p.name.startswith('~$')]
    
    print(f"Scanning {len(team_presentations)} presentations for '{month}' goals...\n")
    
    all_goals = []
    for ppt_file in team_presentations:
        goals = find_and_extract_goals(str(ppt_file), month)
        if goals:
            print(f"✓ {ppt_file.name}: Found {len(goals)} goal slide(s)")
            all_goals.extend(goals)
    
    print(f"\nTotal goal slides found: {len(all_goals)}\n")
    
    # Create Excel report
    report_data = []
    for goal in all_goals:
        row = {
            'Source File': goal['source_file'],
            'Slide Number': goal['slide_number'],
            'Month': goal['month'],
            'Title': goal['title'],
            'Text Content': '\n\n'.join(goal['text_content']),
            'Has Tables': 'Yes' if goal['table_data'] else 'No',
            'Has Images': 'Yes' if goal['has_images'] else 'No',
            'Has Charts': 'Yes' if goal['has_charts'] else 'No',
            'Total Shapes': goal['shape_count']
        }
        report_data.append(row)
    
    df = pd.DataFrame(report_data)
    df.to_excel(output_excel, index=False, sheet_name=f"{month} Goals")
    print(f"✓ Excel report saved: {output_excel}")
    
    # Create manual instructions
    instructions = [
        f"GOALS ACHIEVED IN {month.upper()} - CONSOLIDATION GUIDE",
        "=" * 80,
        "",
        "SUMMARY:",
        f"  Total goal slides found: {len(all_goals)}",
        "",
        "SLIDES TO EXTRACT:",
        ""
    ]
    
    for i, goal in enumerate(all_goals, 1):
        instructions.extend([
            f"{i}. File: {goal['source_file']}",
            f"   Slide: #{goal['slide_number']}",
            f"   Title: {goal['title']}",
            ""
        ])
    
    instructions.extend([
        "",
        "MANUAL CONSOLIDATION STEPS:",
        "=" * 80,
        "",
        "1. Open the template PowerPoint:",
        "   sample_data/final-data/06-01-2026.pptx",
        "",
        "2. For each slide listed above:",
        "   a. Open the source presentation file",
        f"   b. Navigate to the specified slide number",
        "   c. Right-click the slide in the thumbnail pane",
        "   d. Select 'Copy'",
        "   e. Go back to the template presentation",
        "   f. Right-click after slide 5 (or desired position)",
        "   g. Select 'Keep Source Formatting' when pasting",
        "",
        "3. Save the consolidated presentation as:",
        "   output/consolidated_mbr_manual.pptx",
        "",
        "=" * 80,
        "",
        "NOTE: Python-pptx library has limitations with complex slide copying.",
        "Manual copying preserves all formatting, images, and charts correctly.",
        "",
        f"Detailed content extracted to: {Path(output_excel).name}",
        ""
    ])
    
    instructions_text = "\n".join(instructions)
    Path(output_instructions).write_text(instructions_text, encoding='utf-8')
    print(f"✓ Instructions saved: {output_instructions}")
    
    return {
        'goals_found': len(all_goals),
        'excel_report': output_excel,
        'instructions': output_instructions,
        'goals_data': all_goals
    }


if __name__ == "__main__":
    result = create_goals_report(
        ppt_data_folder="sample_data/ppt-data",
        output_excel="output/goals_achieved_report.xlsx",
        output_instructions="output/ppt_consolidation_instructions.txt",
        month="May"
    )
    
    print(f"\n{'=' * 80}")
    print("SUMMARY:")
    print(f"  Goals found: {result['goals_found']}")
    print(f"  Excel report: {result['excel_report']}")
    print(f"  Instructions: {result['instructions']}")
    print("=" * 80)
